import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), "../../data/kp_template.docx")

def ensure_kp_template_exists():
    """Создает базовый шаблон Ростелекома, если файл отсутствует."""
    if os.path.exists(TEMPLATE_PATH):
        return

    os.makedirs(os.path.dirname(TEMPLATE_PATH), exist_ok=True)
    doc = Document()

    # Заголовок
    title = doc.add_heading(level=1)
    run = title.add_run("ПАО «РОСТЕЛЕКОМ» | КОММЕРЧЕСКОЕ ПРЕДЛОЖЕНИЕ")
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(0, 51, 153) # Корпоративный синий
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("Дата формирования: {{ date }}")
    doc.add_paragraph("Клиент / Объект: {{ client_name }} (ИНН: {{ inn }})")
    doc.add_paragraph("Проект развития: {{ project_title }}")

    doc.add_heading("1. Параметры цифрового решения", level=2)
    p = doc.add_paragraph(
        "На основе аудита инфраструктуры экспертами ПАО «Ростелеком» "
        "для вашего объекта сформирована спецификация оборудования и каналов связи:"
    )

    # Таблица со сметой (плейсхолдеры docxtpl)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Наименование работ / Сервис"
    hdr_cells[1].text = "Объем / Масштаб"
    hdr_cells[2].text = "Плановая стоимость (руб)"

    # Строка цикла Jinja
    row_cells = table.add_row().cells
    row_cells[0].text = "{% for item in items %}{{ item.name }}"
    row_cells[1].text = "{{ item.scale }}"
    row_cells[2].text = "{{ item.price }}{% if not loop.last %}\n{% endif %}{% endfor %}"

    doc.add_heading("2. Ценовое предложение и условия", level=2)
    doc.add_paragraph("• Базовая себестоимость решения: {{ cost_price }} руб.")
    doc.add_paragraph("• Оптимальная коммерческая стоимость: {{ optimal_price }} руб.")
    doc.add_paragraph("• Минимальный согласованный порог: {{ min_price }} руб.")
    doc.add_paragraph("• Расчетный уровень маржинальности: {{ margin_percent }}%")

    doc.add_paragraph("\nПредложение действительно в течение 30 календарных дней.")
    doc.add_paragraph("Контакты куратора: B2B Департамент ПАО «Ростелеком», b2b@rt.ru")

    doc.save(TEMPLATE_PATH)
    print(f"[+] Сгенерирован корпоративный шаблон КП: {TEMPLATE_PATH}")